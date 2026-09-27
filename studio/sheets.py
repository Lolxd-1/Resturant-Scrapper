"""SmartBiz output builder: upload-ready sheet, pending sheet, and a photo report workbook."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from studio.models import DishResult

STATUS_LABELS = {
    "ready": "Ready",
    "needs_photo": "Needs owner photo",
    "needs_price": "Needs price",
    "error": "Error",
}

DISHES_HEADER = [
    "SKU", "Dish", "Category", "Price", "Status", "Reason", "Confidence",
    "Match", "Source", "Zomato item", "Restaurant", "Restaurant link", "Photo link",
]
DISHES_WIDTHS = [14, 38, 16, 12, 16, 40, 12, 12, 12, 18, 18, 30, 30]

_TEMPLATE_COLUMNS = (3, 4, 5, 6, 7, 8, 9, 16)  # C,D,E,F,G,H,I,P

CHECKLIST_HEADER = ["SKU", "Dish", "Category", "Why", "How to shoot"]
HOW_TO_SHOOT = (
    "Top-down or 45° shot in daylight, dish centred on a plain plate, "
    "no text or logos, at least 800x800 px."
)

_SAFE_RE = re.compile(r"[^A-Za-z0-9_]+")


def safe_vendor(vendor: str) -> str:
    cleaned = _SAFE_RE.sub("_", vendor or "").strip("_")
    cleaned = re.sub(r"_+", "_", cleaned)
    return cleaned[:40] or "Vendor"


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _price_value(price: float | None):
    if price is None:
        return None
    return int(price) if float(price).is_integer() else round(price, 2)


def category_for(r: DishResult) -> str:
    cat = (r.category or "").strip()
    return cat if cat else "Menu"


def _write_rows(ws: Worksheet, rows: list[DishResult]) -> None:
    for i, r in enumerate(rows, start=2):
        price = _price_value(r.price)
        ws.cell(row=i, column=3).value = _clean(r.sku)
        ws.cell(row=i, column=4).value = _clean(r.name[:200])
        ws.cell(row=i, column=5).value = price
        ws.cell(row=i, column=6).value = price
        ws.cell(row=i, column=7).value = "FOOD_AND_GROCERY"
        ws.cell(row=i, column=8).value = _clean(category_for(r)[:100])
        ws.cell(row=i, column=9).value = _clean(r.description[:2000]) if r.description else None
        ws.cell(row=i, column=16).value = _clean(r.hosted_url) if r.hosted_url else None


def _clear_rows(ws: Worksheet, count: int) -> None:
    """Reset exactly the cells a prior `_write_rows(ws, rows[:count])` touched."""
    for i in range(2, count + 2):
        for col in _TEMPLATE_COLUMNS:
            ws.cell(row=i, column=col).value = None


def _set_link(cell, url: str | None) -> None:
    if url:
        url = _clean(url)
        cell.value = url
        cell.hyperlink = url


def _write_summary(ws: Worksheet, results: list[DishResult], vendor: str, vision_on: bool,
                   logo_on: bool) -> None:
    total = len(results)
    ready_n = sum(1 for r in results if r.status == "ready")
    needs_photo_n = sum(1 for r in results if r.status == "needs_photo")
    needs_price_n = sum(1 for r in results if r.status == "needs_price")
    error_n = sum(1 for r in results if r.status == "error")
    needs_price_with_photo = sum(1 for r in results if r.status == "needs_price" and r.hosted_url)
    coverage = round((ready_n + needs_price_with_photo) / total * 100) if total else 0

    rows = [
        ("Vendor", _clean(vendor)),
        ("Generated", datetime.now().strftime("%d %b %Y %H:%M")),
        ("Dishes in menu", total),
        ("Ready to upload", ready_n),
        ("Need owner photo", needs_photo_n),
        ("Need price", needs_price_n),
        ("Errors", error_n),
        ("Photo coverage", f"{coverage}%"),
        ("Visual check", "On" if vision_on else "Unavailable"),
        ("Logo check", "On" if logo_on else "Unavailable"),
    ]
    for label, value in rows:
        ws.append([label, value])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 30


def _write_dishes(ws: Worksheet, results: list[DishResult]) -> None:
    ws.append(DISHES_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in results:
        row = [
            r.sku,
            _clean(r.name),
            _clean(category_for(r)),
            _price_value(r.price),
            STATUS_LABELS.get(r.status, r.status),
            _clean(r.reason),
            round(r.confidence * 100),
            r.tier,
            r.source,
            _clean(r.item_name),
            _clean(r.restaurant),
            None,
            None,
        ]
        ws.append(row)
        out_row = ws.max_row
        _set_link(ws.cell(row=out_row, column=12), r.restaurant_url)
        _set_link(ws.cell(row=out_row, column=13), r.hosted_url)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DISHES_HEADER))}{ws.max_row}"
    for idx, width in enumerate(DISHES_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_checklist(ws: Worksheet, results: list[DishResult]) -> None:
    ws.append(CHECKLIST_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in results:
        if r.status != "needs_photo":
            continue
        ws.append([r.sku, _clean(r.name), _clean(category_for(r)), _clean(r.reason), HOW_TO_SHOOT])
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 60


def _build_report(path: Path, results: list[DishResult], vendor: str, vision_on: bool, logo_on: bool) -> None:
    wb = openpyxl.Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    _write_summary(summary_ws, results, vendor, vision_on, logo_on)
    _write_dishes(wb.create_sheet("Dishes"), results)
    _write_checklist(wb.create_sheet("Owner photo checklist"), results)
    wb.save(path)
    wb.close()


def build_outputs(
    results: list[DishResult], template: Path, out_dir: Path, vendor: str, *, vision_on: bool = True,
    logo_on: bool = True
) -> dict[str, Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    ordered = sorted(results, key=lambda r: r.row)
    ready = [r for r in ordered if r.status == "ready"]
    pending = [r for r in ordered if r.status != "ready"]

    stem = safe_vendor(vendor)
    ready_path = out_dir / f"{stem}_SmartBiz_UPLOAD_READY.xlsx"
    pending_path = out_dir / f"{stem}_SmartBiz_PENDING.xlsx"
    report_path = out_dir / f"{stem}_Photo_Report.xlsx"

    wb = openpyxl.load_workbook(template)
    ws = wb["bulk_upload_template"]
    _write_rows(ws, ready)
    wb.save(ready_path)
    _clear_rows(ws, len(ready))
    _write_rows(ws, pending)
    wb.save(pending_path)
    wb.close()

    _build_report(report_path, ordered, vendor, vision_on, logo_on)

    return {"ready": ready_path, "pending": pending_path, "report": report_path}
