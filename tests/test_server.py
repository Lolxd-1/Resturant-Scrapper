"""Tests for studio.server (FastAPI app, background worker, SSE) — offline (FakeSession HTTP)."""
from __future__ import annotations

import io
import json
import time
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

from studio import pipeline
from studio import server as server_mod
from studio.models import PoolItem, RestaurantMenu
from studio.pipeline import Engine
from studio.store import PoolStore
from studio.zomato import Http
from tests.conftest import FakeResponse, FakeSession, make_xlsx
from tests.test_pipeline import FoodVision, noise_jpeg
from tests.test_sheets import make_template

IMG = "https://b.zmtcdn.com/data/dish/"
RES_URL = "https://www.zomato.com/pune/pizza-place"


@pytest.fixture(autouse=True)
def small_template(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "TEMPLATE_PATH", make_template(tmp_path / "template.xlsx"))


def _build_engine(tmp_path) -> Engine:
    session = FakeSession()
    http = Http(session=session, min_interval=0.0, sleep=lambda s: None)
    url = IMG + "cheese1.jpg"
    session.add(url, FakeResponse(200, content=noise_jpeg(1)))
    store = PoolStore(tmp_path / "pool.sqlite")
    items = [PoolItem(0, "Cheese Pizza", "Pizza", "Delivery", "", "veg", url,
                      "r1", "Pizza Place", RES_URL, "pune", "")]
    store.upsert_menu(RestaurantMenu("r1", "Pizza Place", RES_URL, "pune", "", items))
    return Engine(store, FoodVision(), http)


def make_app(tmp_path, start_worker=True):
    engine = _build_engine(tmp_path)
    output_dir = tmp_path / "outputs"
    config_path = tmp_path / "config" / "settings.json"

    def factory():
        return engine

    app = server_mod.create_app(engine_factory=factory, output_dir=output_dir,
                                config_path=config_path, start_worker=start_worker)
    return app, output_dir, config_path


def wait_for_engine(app, timeout=5.0):
    worker = app.state.worker
    deadline = time.time() + timeout
    while worker.engine is None and time.time() < deadline:
        time.sleep(0.02)
    assert worker.engine is not None


def upload_menu(client, tmp_path, rows, name="menu.xlsx"):
    path = make_xlsx(tmp_path / name, [["Item Name", "Category", "Description", "Price"], *rows])
    with open(path, "rb") as f:
        return client.post(
            "/api/menu/preview",
            files={"file": (name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )


def _disable_live_search(client):
    r = client.put("/api/settings", json={
        "imgbb_keys": [], "strictness": "balanced", "live_search": False,
        "live_budget_menus": 0, "default_city": "pune",
    })
    assert r.status_code == 200


def _run_simple_job(client, tmp_path, vendor="Test Cafe"):
    _disable_live_search(client)
    r = upload_menu(client, tmp_path, [["Cheese Pizza", "Pizza", "", 199]])
    assert r.status_code == 200
    upload_id = r.json()["upload_id"]
    r2 = client.post("/api/jobs", json={"upload_id": upload_id, "vendor": vendor, "city": "pune"})
    assert r2.status_code == 200
    job_id = r2.json()["id"]

    deadline = time.time() + 10
    job = None
    while time.time() < deadline:
        jr = client.get(f"/api/jobs/{job_id}")
        job = jr.json()
        if job["job"]["status"] in ("done", "failed", "cancelled"):
            break
        time.sleep(0.05)
    assert job is not None and job["job"]["status"] == "done", job
    return job_id, job


# ---- TC1 ----

def test_health_and_index(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)

    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json() == {"ok": True, "app": "menustudio", "version": server_mod.studio.__version__}

    r2 = client.get("/")
    assert r2.status_code == 200
    assert "Menu Studio" in r2.text

    r3 = client.get("/api/system")
    assert r3.status_code == 200
    assert isinstance(r3.json()["free_gb"], float)


# ---- TC2 ----

def test_settings_roundtrip_and_validation(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    key1 = "a" * 32
    key2 = "b" * 32

    r = client.put("/api/settings", json={
        "imgbb_keys": [{"label": "One", "key": key1, "enabled": True},
                       {"label": "Two", "key": key2, "enabled": True}],
        "strictness": "balanced", "live_search": True, "live_budget_menus": 100, "default_city": "pune",
    })
    assert r.status_code == 200
    assert key1 not in r.text and key2 not in r.text
    body = r.json()
    assert [k["tail"] for k in body["imgbb_keys"]] == [key1[-4:], key2[-4:]]

    r2 = client.get("/api/settings")
    assert key1 not in r2.text and key2 not in r2.text
    assert sorted(r2.json()["cities"]) == r2.json()["cities"]

    r3 = client.put("/api/settings", json={"imgbb_keys": [{"id": 0, "label": "Main"}]})
    assert r3.status_code == 200
    assert r3.json()["imgbb_keys"][0]["tail"] == key1[-4:]
    assert r3.json()["imgbb_keys"][0]["label"] == "Main"

    r4 = client.put("/api/settings", json={"imgbb_keys": [{"label": "Bad", "key": "short"}]})
    assert r4.status_code == 400
    assert r4.json()["detail"] == "That doesn't look like an ImgBB key"


# ---- TC3 ----

def test_menu_preview(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)

    r = upload_menu(client, tmp_path, [["Cheese Pizza", "Pizza", "", 199], ["Cold Coffee", "Beverages", "", 120]])
    assert r.status_code == 200
    data = r.json()
    assert data["rows"][0]["name"] == "Cheese Pizza"
    assert {c["name"] for c in data["categories"]} == {"Pizza", "Beverages"}
    assert data["upload_id"]

    xls_path = tmp_path / "menu.xls"
    xls_path.write_bytes(b"not really an xls file")
    with open(xls_path, "rb") as f:
        r2 = client.post("/api/menu/preview", files={"file": ("menu.xls", f, "application/vnd.ms-excel")})
    assert r2.status_code == 200
    assert r2.json()["errors"]

    big = b"0" * (11 * 1024 * 1024)
    r3 = client.post("/api/menu/preview", files={"file": ("big.xlsx", io.BytesIO(big), "application/octet-stream")})
    assert r3.status_code == 413


# ---- TC4 ----

def test_job_flow_files_and_zip(tmp_path):
    app, output_dir, config_path = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)

    job_id, job = _run_simple_job(client, tmp_path)
    assert job["results"]

    r = client.get(f"/api/jobs/{job_id}/files/ready")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"

    r2 = client.get(f"/api/jobs/{job_id}/zip")
    assert r2.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(r2.content)).namelist()
    assert sum(1 for n in names if n.endswith(".xlsx")) == 3


# ---- TC5 ----

def test_events_replay_includes_done_and_closes(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)

    job_id, _job = _run_simple_job(client, tmp_path)

    r = client.get(f"/api/jobs/{job_id}/events?after=-1")
    assert r.status_code == 200
    events = []
    for block in r.text.split("\n\n"):
        for line in block.split("\n"):
            if line.startswith("data:"):
                events.append(json.loads(line[len("data:"):].strip()))
    assert any(e["type"] == "done" for e in events)


# ---- TC6 ----

def test_validation_404_and_400(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)

    r = client.get("/api/jobs/../../etc/files/ready")
    assert r.status_code == 404

    job_id, _job = _run_simple_job(client, tmp_path)
    r2 = client.get(f"/api/jobs/{job_id}/files/unknownkind")
    assert r2.status_code == 404

    r3 = client.post("/api/jobs", json={"vendor": "", "upload_id": "x"})
    assert r3.status_code == 400


# ---- TC7 ----

# ---- FIX2 TC6 ----

def test_fix2_tc6_zomato_url_validation(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    _disable_live_search(client)

    r = upload_menu(client, tmp_path, [["Cheese Pizza", "Pizza", "", 199]])
    upload_id = r.json()["upload_id"]

    bad = client.post("/api/jobs", json={
        "upload_id": upload_id, "vendor": "Test Cafe", "city": "pune", "zomato_url": "zoma.to/abc",
    })
    assert bad.status_code == 400
    assert "full zomato.com address" in bad.json()["detail"]

    r2 = upload_menu(client, tmp_path, [["Cheese Pizza", "Pizza", "", 199]], name="menu2.xlsx")
    upload_id2 = r2.json()["upload_id"]
    good = client.post("/api/jobs", json={
        "upload_id": upload_id2, "vendor": "Test Cafe", "city": "pune",
        "zomato_url": "www.zomato.com/pune/x",
    })
    assert good.status_code == 200


def test_choose_null_marks_needs_photo(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)

    job_id, job = _run_simple_job(client, tmp_path)
    sku = job["results"][0]["sku"]

    r = client.post(f"/api/jobs/{job_id}/dishes/{sku}/choose", json={"alt_index": None})
    assert r.status_code == 200
    assert r.json()["status"] == "needs_photo"


# ---- FIX1 ----

def _two_ready_dishes(output_dir, job_id) -> tuple[str, str]:
    """Rewrite the job's results as two ready dishes with hosted photos; return their SKUs."""
    results_file = output_dir / job_id / "results.json"
    data = json.loads(results_file.read_text(encoding="utf-8"))
    first = data[0]
    sku_a = first["sku"]
    sku_b = sku_a.rsplit("-", 1)[0] + "-099"
    first.update(status="ready", reason="", hosted_url="https://i.ibb.co/1/a.jpg", zomato_img=IMG + "cheese1.jpg")
    second = dict(first, sku=sku_b, name="Paneer Pizza", hosted_url="https://i.ibb.co/2/b.jpg")
    results_file.write_text(json.dumps([first, second]), encoding="utf-8")
    return sku_a, sku_b


def _ready_rows(content: bytes) -> list:
    ws = openpyxl.load_workbook(io.BytesIO(content))["bulk_upload_template"]
    return [ws[f"D{r}"].value for r in range(2, 5) if ws[f"D{r}"].value]


def test_fix1_tc2_choose_while_job_active_is_409(tmp_path, monkeypatch):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    job_id, job = _run_simple_job(client, tmp_path)
    monkeypatch.setattr(app.state.worker, "is_active", lambda jid: True)

    r = client.post(f"/api/jobs/{job_id}/dishes/{job['results'][0]['sku']}/choose", json={"alt_index": None})

    assert r.status_code == 409
    assert r.json()["detail"] == "This job is still running — wait until it finishes, then change photos."


def test_fix1_tc3_download_waits_for_latest_swap(tmp_path, monkeypatch):
    app, output_dir, _ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    job_id, _job = _run_simple_job(client, tmp_path)
    sku_a, sku_b = _two_ready_dishes(output_dir, job_id)
    real_rebuild = server_mod.rebuild_outputs

    def slow_rebuild(engine, job_dir):
        summary = real_rebuild(engine, job_dir)
        time.sleep(0.5)
        return summary

    monkeypatch.setattr(server_mod, "rebuild_outputs", slow_rebuild)

    ra = client.post(f"/api/jobs/{job_id}/dishes/{sku_a}/choose", json={"alt_index": None})
    rb = client.post(f"/api/jobs/{job_id}/dishes/{sku_b}/choose", json={"alt_index": None})
    assert ra.status_code == 200 and rb.status_code == 200

    r = client.get(f"/api/jobs/{job_id}/files/ready")

    assert r.status_code == 200
    assert _ready_rows(r.content) == []
    job = json.loads((output_dir / job_id / "job.json").read_text(encoding="utf-8"))
    assert job["sheets_stale"] is False


def test_fix1_tc4_failed_rebuild_is_409_until_fixed(tmp_path, monkeypatch):
    app, output_dir, _ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    job_id, _job = _run_simple_job(client, tmp_path)
    sku_a, _sku_b = _two_ready_dishes(output_dir, job_id)
    real_rebuild = server_mod.rebuild_outputs

    def locked_rebuild(engine, job_dir):
        raise PermissionError(13, "Permission denied", str(job_dir / "Test_Cafe_SmartBiz_UPLOAD_READY.xlsx"))

    monkeypatch.setattr(server_mod, "rebuild_outputs", locked_rebuild)
    assert client.post(f"/api/jobs/{job_id}/dishes/{sku_a}/choose", json={"alt_index": None}).status_code == 200

    r = client.get(f"/api/jobs/{job_id}/files/ready")
    assert r.status_code == 409
    assert r.json()["detail"] == "Close the Excel file (Test_Cafe_SmartBiz_UPLOAD_READY.xlsx) and try again."
    assert client.get(f"/api/jobs/{job_id}/zip").status_code == 409

    monkeypatch.setattr(server_mod, "rebuild_outputs", real_rebuild)
    r2 = client.get(f"/api/jobs/{job_id}/files/ready")
    assert r2.status_code == 200
    assert _ready_rows(r2.content) == ["Paneer Pizza"]
    job = json.loads((output_dir / job_id / "job.json").read_text(encoding="utf-8"))
    assert "rebuild_error" not in job and job["sheets_stale"] is False


# ---- TC8 ----

def test_app_self_test():
    import app as app_module
    assert app_module.main(["--self-test"]) == 0


# ---- A6: search + approve routes ----

def _write_results(output_dir, job_id, rows):
    (output_dir / job_id / "results.json").write_text(json.dumps(rows), encoding="utf-8")


def test_a6_tc1_search_route(tmp_path, monkeypatch):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    job_id, job = _run_simple_job(client, tmp_path)
    row = job["results"][0]
    calls = []

    def fake_search(engine, job_dir, sku, query):
        calls.append((sku, query))
        return pipeline.dish_result_from_dict(dict(row, alternates=[{"item_name": "A"}, {"item_name": "B"}]))

    monkeypatch.setattr(server_mod, "search_photos", fake_search)

    r = client.post(f"/api/jobs/{job_id}/dishes/{row['sku']}/search", json={"query": "Veg Thali"})

    assert r.status_code == 200
    assert len(r.json()["alternates"]) == 2
    assert calls == [(row["sku"], "Veg Thali")]


def test_a6_tc2_search_guards(tmp_path, monkeypatch):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    job_id, job = _run_simple_job(client, tmp_path)
    sku = job["results"][0]["sku"]

    empty = client.post(f"/api/jobs/{job_id}/dishes/{sku}/search", json={"query": "   "})
    assert empty.status_code == 400 and empty.json()["detail"] == "Type a dish name to search for"
    missing = client.post(f"/api/jobs/20200101-000000-nope/dishes/{sku}/search", json={"query": "x"})
    assert missing.status_code == 404
    monkeypatch.setattr(app.state.worker, "is_active", lambda jid: True)
    busy = client.post(f"/api/jobs/{job_id}/dishes/{sku}/search", json={"query": "x"})
    assert busy.status_code == 409


def test_a6_tc3_approve_all_takes_review_dishes_and_rebuilds(tmp_path, monkeypatch):
    app, output_dir, _ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    job_id, job = _run_simple_job(client, tmp_path)
    base = job["results"][0]
    prefix = base["sku"].rsplit("-", 1)[0]
    skus = [f"{prefix}-001", f"{prefix}-002", f"{prefix}-003"]
    _write_results(output_dir, job_id, [
        dict(base, sku=skus[0], status="review"),
        dict(base, sku=skus[1], status="ready"),
        dict(base, sku=skus[2], status="review"),
    ])
    got = []
    monkeypatch.setattr(server_mod, "approve_dishes", lambda engine, job_dir, s, uploader: got.append(s) or [])
    gen_before = app.state.worker._gen.get(job_id, 0)

    r = client.post(f"/api/jobs/{job_id}/approve", json={"all": True})

    assert r.status_code == 200 and r.json() == {"results": []}
    assert got == [[skus[0], skus[2]]]
    assert app.state.worker._gen[job_id] == gen_before + 1


def test_a6_tc4_approve_rejects_bad_sku(tmp_path):
    app, *_ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    job_id, _job = _run_simple_job(client, tmp_path)

    r = client.post(f"/api/jobs/{job_id}/approve", json={"skus": ["../x"]})

    assert r.status_code == 400


def test_a6_tc5_approve_real_lookalike_becomes_ready(tmp_path):
    app, output_dir, _ = make_app(tmp_path)
    client = TestClient(app)
    wait_for_engine(app)
    _disable_live_search(client)
    url = IMG + "cheese1.jpg"
    # retrieval skips tokens found in over a quarter of a tiny pool, so pad it
    names = [("Paneer Pizza", "Pizza", "veg"), ("Cold Coffee", "Drinks", "veg"),
             ("Mango Shake", "Shakes", "veg"), ("Veg Burger", "Burgers", "veg"),
             ("Masala Dosa", "South", "veg"), ("Fries", "Sides", "veg"), ("Brownie", "Desserts", "veg"),
             ("Dal Rice", "Meals", "veg")]
    items = [PoolItem(0, n, cat, "Delivery", "", diet, url, "r2", "Paneer Place", RES_URL + "2", "pune", "")
             for n, cat, diet in names]
    app.state.worker.engine.add_items(RestaurantMenu("r2", "Paneer Place", RES_URL + "2", "pune", "", items))
    upload_id = upload_menu(client, tmp_path, [["Green Chilli Paneer Pizza (Single Cheese)", "Pizza", "", 120]]).json()["upload_id"]
    job_id = client.post("/api/jobs", json={"upload_id": upload_id, "vendor": "Test Cafe", "city": "pune"}).json()["id"]
    deadline = time.time() + 10
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["job"]["status"] in ("done", "failed", "cancelled"):
            break
        time.sleep(0.05)
    dish = job["results"][0]
    assert dish["status"] == "review", dish
    job_file = output_dir / job_id / "job.json"
    meta = json.loads(job_file.read_text(encoding="utf-8"))
    meta["options"]["upload"] = False
    job_file.write_text(json.dumps(meta), encoding="utf-8")

    r = client.post(f"/api/jobs/{job_id}/approve", json={"skus": [dish["sku"]]})

    assert r.status_code == 200
    assert [d["status"] for d in r.json()["results"]] == ["ready"]
