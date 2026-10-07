"""Offline end-to-end tests for studio.pipeline (FakeSession HTTP, tmp pool, fake vision, fake ImgBB)."""
from __future__ import annotations

import io
import json
import sys
import threading
import zlib
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np
import openpyxl
import pytest
from PIL import Image, ImageDraw, ImageFont

from studio import pipeline
from studio.brand import BrandGuard
from studio.config import atomic_write_text
from studio.hosting import Uploader
from studio.imaging import THUMB_SUFFIX, full_url, thumb_url
from studio.models import Candidate, DishQuery, ImgbbKey, MatchScore, PoolItem, RestaurantMenu
from studio.pipeline import (
    NO_KEY_MESSAGE,
    SWIGGY_BLOCKED_MESSAGE,
    Engine,
    JobOptions,
    approve_dishes,
    choose_alternate,
    list_jobs,
    load_job,
    rebuild_outputs,
    run_job,
    search_photos,
)
from studio.sheets import safe_vendor
from studio.store import PoolStore
from studio.swiggy import SEARCH_URL, SwiggyClient, parse_dish_search
from studio.vision import FOOD_PROMPTS, NullVision
from studio.zomato import Http
from tests.conftest import FakeResponse, FakeSession, make_xlsx
from tests.fixtures.zomato.make_pages import item, menu_state, search_state, state_html, suggest_json
from tests.test_sheets import make_template

IMG = "https://b.zmtcdn.com/data/dish/"
SUGGEST = "https://www.zomato.com/webroutes/search/autoSuggest"
HEADER = ["Item Name", "Category", "Description", "Price"]


class FoodVision:
    """Every image embeds to the "food on a plate" prompt; texts get distinct one-hot vectors."""

    available = True
    reason = ""

    def __init__(self):
        self._index: dict[str, int] = {}

    def embed_texts(self, texts):
        vecs = np.zeros((len(texts), 512), dtype=np.float32)
        for row, t in enumerate(texts):
            vecs[row, self._index.setdefault(t, len(self._index))] = 1.0
        return vecs

    def embed_images(self, images):
        return np.repeat(self.embed_texts([FOOD_PROMPTS[0]]), len(images), axis=0)


def noise_jpeg(seed: int) -> bytes:
    blocks = np.random.default_rng(seed).integers(0, 256, (6, 6, 3), dtype=np.uint8)
    buf = io.BytesIO()
    Image.fromarray(blocks).resize((420, 420), Image.NEAREST).save(buf, format="JPEG", quality=80)
    return buf.getvalue()


class Env:
    def __init__(self, tmp_path, vision=None):
        self.tmp = tmp_path
        self.session = FakeSession()
        self.http = Http(session=self.session, min_interval=0.0, sleep=lambda s: None)
        self.store = PoolStore(tmp_path / "pool.sqlite")
        self.vision = vision or FoodVision()
        self.posts: list[dict] = []
        self.uploader = Uploader([ImgbbKey("t", "k" * 32)], http_post=self.post, store=self.store)
        self.events: list[dict] = []

    def post(self, url, data=None, files=None, timeout=None):
        self.posts.append(data)
        n = len(self.posts)
        body = {"success": True, "data": {"url": f"https://i.ibb.co/{n}/x.jpg",
                                          "display_url": f"https://i.ibb.co/{n}/x.jpg"}}
        return FakeResponse(200, json.dumps(body))

    def image(self, name: str) -> str:
        url = IMG + name + ".jpg"
        if url not in self.session.routes:
            self.session.add(url, FakeResponse(200, content=noise_jpeg(zlib.crc32(name.encode()))))
        return url

    def pool(self, slug: str, dishes, city: str = "pune", res_name: str | None = None) -> None:
        """dishes: (name, category, diet, image name)."""
        url = f"https://www.zomato.com/{city}/{slug}"
        name = res_name or slug.replace("-", " ").title()
        items = [PoolItem(0, n, cat, "Delivery", "", diet, self.image(img), slug, name, url, city, "")
                 for n, cat, diet, img in dishes]
        self.store.upsert_menu(RestaurantMenu(slug, name, url, city, "", items))

    def page(self, slug: str, dishes, city: str = "pune", res_name: str | None = None) -> None:
        """A live Zomato restaurant page (not in the pool)."""
        url = f"https://www.zomato.com/{city}/{slug}"
        cats: dict[str, list] = {}
        for n, cat, diet, img in dishes:
            cats.setdefault(cat, []).append(item(n, self.image(img), diet))
        state = menu_state(slug, res_name or slug.replace("-", " ").title(), url,
                           [("Delivery", list(cats.items()))])
        self.session.add(url + "/order", FakeResponse(200, state_html(state)))

    def engine(self) -> Engine:
        return Engine(self.store, self.vision, self.http)

    def menu(self, rows, header=HEADER, name="menu.xlsx"):
        return make_xlsx(self.tmp / name, [header, *rows])

    def run(self, menu_path, opts, job="job1", engine=None, upload=True, **kw):
        opts.upload = upload
        kw.setdefault("uploader", self.uploader if upload else None)
        kw.setdefault("on_event", self.events.append)
        return run_job(engine or self.engine(), menu_path, self.tmp / job, opts, **kw)

    def results(self, job="job1") -> dict[str, dict]:
        data = json.loads((self.tmp / job / "results.json").read_text(encoding="utf-8"))
        return {r["name"]: r for r in data}


@pytest.fixture(autouse=True)
def small_template(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "TEMPLATE_PATH", make_template(tmp_path / "template.xlsx"))


def sheet(env, vendor, kind, job="job1"):
    return openpyxl.load_workbook(env.tmp / job / f"{safe_vendor(vendor)}_SmartBiz_{kind}.xlsx")["bulk_upload_template"]


def col(ws, letter, n):
    return [ws[f"{letter}{r}"].value for r in range(2, n + 2)]


# ---- TC1: end to end with a seeded pool ----

def test_end_to_end_ready_pending_images(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizzas", "veg", "cheese1"),
                             ("Farmhouse Pizza", "Pizzas", "veg", "farm1"),
                             ("Cold Coffee", "Beverages", "veg", "cc1"),
                             ("Mango Shake", "Shakes", "veg", "mango1")])
    env.pool("cafe-two", [("Cheese Pizza", "Pizza", "veg", "cheese2"),
                          ("Cold Coffee", "Drinks", "veg", "cc2")])
    menu = env.menu([
        ["Cheese Pizza - Small", "Pizza", "", 199],
        ["Cheese Pizza - Large", "Pizza", "", 399],
        ["Farmhouse Pizza", "Pizza", "", 299],
        ["Cold Coffee", "Beverages", "", 120],
        ["Mango Shake", "Beverages", "Thick mango shake", None],
        ["Ceet-B Mastani", "Beverages", "", 90],
    ])
    summary = env.run(menu, JobOptions(vendor="Test Cafe", live_search=False))

    assert summary["total"] == 6 and summary["ready"] == 4 and summary["needs_price"] == 1
    ready = sheet(env, "Test Cafe", "UPLOAD_READY")
    assert col(ready, "D", 5) == ["Cheese Pizza - Small", "Cheese Pizza - Large", "Farmhouse Pizza",
                                  "Cold Coffee", None]
    assert col(ready, "H", 4) == ["Pizza", "Pizza", "Pizza", "Beverages"]
    assert all(u and u.startswith("https://i.ibb.co/") for u in col(ready, "P", 4))
    pending = sheet(env, "Test Cafe", "PENDING")
    assert col(pending, "D", 2) == ["Mango Shake", "Ceet-B Mastani"]
    assert pending["P2"].value.startswith("https://i.ibb.co/")

    res = env.results()
    small, large = res["Cheese Pizza - Small"], res["Cheese Pizza - Large"]
    assert small["zomato_img"] and small["zomato_img"] == large["zomato_img"]
    assert small["hosted_url"] == large["hosted_url"]
    for r in res.values():
        if r["local_image"]:
            assert (tmp_path / "job1" / r["local_image"]).is_file()
    assert (tmp_path / "job1" / "images" / "TC-001.jpg").is_file()
    assert res["Ceet-B Mastani"]["status"] == "needs_photo"
    assert res["Ceet-B Mastani"]["reason"] == "Not found on Zomato"
    job = load_job(tmp_path / "job1")
    assert job["status"] == "done" and job["summary"]["ready"] == 4 and job["menu_file"] == "menu.xlsx"
    assert list_jobs(tmp_path)[0]["id"] == "job1"
    assert any(e["type"] == "done" for e in env.events)


# ---- TC2: vendor is a chain on Zomato ----

def test_chain_mode_uses_chain_photos(tmp_path):
    env = Env(tmp_path)
    env.pool("other-place", [("Farmhouse Pizza", "Pizza", "veg", "farm-other"),
                             ("Cold Coffee", "Drinks", "veg", "cc-other")])
    env.session.add(SUGGEST, FakeResponse(200, suggest_json(
        [{"entityType": "chain_outlets", "actionInfo": {"clickUrl": "/pune/brands/pizzaco"}}])))
    env.session.add("https://www.zomato.com/pune/brands/pizzaco", FakeResponse(200, state_html(
        search_state(["/pune/pizzaco-baner/order", "/pune/pizzaco-kothrud/order"]))))
    env.page("pizzaco-baner", [("Farmhouse Pizza", "Pizza", "veg", "farm-chain1")], res_name="PizzaCo")
    env.page("pizzaco-kothrud", [("Farmhouse Pizza", "Pizza", "veg", "farm-chain2")], res_name="PizzaCo")
    menu = env.menu([["Farmhouse Pizza", "Pizza", "", 299], ["Cold Coffee", "Beverages", "", 120]])

    env.run(menu, JobOptions(vendor="PizzaCo", live_search=False), upload=False)

    res = env.results()
    assert res["Farmhouse Pizza"]["source"] == "chain"
    assert res["Farmhouse Pizza"]["zomato_img"] in (IMG + "farm-chain1.jpg", IMG + "farm-chain2.jpg")
    assert res["Cold Coffee"]["source"] == "pool"
    assert res["Farmhouse Pizza"]["status"] == "ready"
    assert any("found on Zomato as chain" in e.get("message", "") for e in env.events)


def test_tune2_tc1_chain_uses_pool_outlets_of_same_brand(tmp_path):
    env = Env(tmp_path)
    env.pool("pizzaco-aundh", [("Garden Special", "Pizzas", "veg", "garden-chain")], res_name="PizzaCo")
    env.pool("other-place", [("Garden Special Pizza", "Pizzas", "veg", "garden-other")])
    env.session.add(SUGGEST, FakeResponse(200, suggest_json(
        [{"entityType": "chain_outlets", "actionInfo": {"clickUrl": "/pune/brands/pizzaco"}}])))
    env.session.add("https://www.zomato.com/pune/brands/pizzaco", FakeResponse(200, state_html(
        search_state(["/pune/pizzaco-baner/order", "/pune/pizzaco-kothrud/order"]))))
    env.page("pizzaco-baner", [("Farmhouse Pizza", "Pizza", "veg", "farm-chain1")], res_name="PizzaCo")
    env.page("pizzaco-kothrud", [("Farmhouse Pizza", "Pizza", "veg", "farm-chain2")], res_name="PizzaCo")
    menu = env.menu([["Garden Special", "Pizzas", "", 299]])

    env.run(menu, JobOptions(vendor="PizzaCo", live_search=False), upload=False)

    r = env.results()["Garden Special"]
    assert r["source"] == "chain" and r["zomato_img"] == IMG + "garden-chain.jpg"
    assert any(e.get("message") == "Using 3 outlets of PizzaCo" for e in env.events)


# ---- TC3: user-supplied own Zomato URL ----

def test_own_restaurant_url(tmp_path):
    env = Env(tmp_path)
    env.pool("other-place", [("Farmhouse Pizza", "Pizza", "veg", "farm-other")])
    env.page("my-own-cafe", [("Farmhouse Pizza", "Pizza", "veg", "farm-own")])
    menu = env.menu([["Farmhouse Pizza", "Pizza", "", 299]])

    env.run(menu, JobOptions(vendor="My Own Cafe", zomato_url="https://www.zomato.com/pune/my-own-cafe/order",
                             live_search=False), upload=False)

    r = env.results()["Farmhouse Pizza"]
    assert r["source"] == "own" and r["zomato_img"] == IMG + "farm-own.jpg" and r["status"] == "ready"
    assert not any("autoSuggest" in url for url, _ in env.session.calls)


# ---- TC4: live whole-Zomato search across cities, persisted in the pool ----

def test_live_search_multi_city_and_persist(tmp_path):
    env = Env(tmp_path)
    env.session.add(SUGGEST, FakeResponse(200, suggest_json([])))
    env.session.add("https://www.zomato.com/mumbai/restaurants/pizza",
                    FakeResponse(200, '<a href="/mumbai/tasty-slice/order">Tasty Slice</a>'))
    env.page("tasty-slice", [("Farmhouse Pizza", "Pizza", "veg", "farm-live")], city="mumbai")
    menu = env.menu([["Farmhouse Pizza", "Pizza", "", 299]])

    env.run(menu, JobOptions(vendor="Tiny Cafe", city="pune"), upload=False)

    r = env.results()["Farmhouse Pizza"]
    assert r["status"] == "ready" and r["source"] == "live"
    assert env.store.menu_fetched_at("https://www.zomato.com/mumbai/tasty-slice") is not None
    assert any(url.startswith("https://www.zomato.com/pune/restaurants/") for url, _ in env.session.calls)

    before = len(env.session.calls)
    env.run(menu, JobOptions(vendor="Tiny Cafe", city="pune"), job="job2", upload=False)
    assert env.results("job2")["Farmhouse Pizza"]["status"] == "ready"
    assert not any(url.endswith("/order") for url, _ in env.session.calls[before:])


# ---- S2 TC1: a dish with no head word crawls the cuisine its words imply ----

def test_s2_tc1_word_driven_cuisine_crawl(tmp_path):
    env = Env(tmp_path)
    env.session.add(SUGGEST, FakeResponse(200, suggest_json([])))
    env.session.add("https://www.zomato.com/pune/restaurants/maharashtrian",
                    FakeResponse(200, '<a href="/pune/kaju-place/order">Kaju Place</a>'))
    env.page("kaju-place", [("Kaju Kolhapuri", "Mains", "veg", "kaju1"),
                            ("Chicken Biryani", "Mains", "nonveg", "cb1"),
                            ("Paneer Tikka", "Starters", "veg", "pt1"),
                            ("Veg Fried Rice", "Rice", "veg", "vfr1")])
    menu = env.menu([["Kaju Kolhapuri", "Mains", "", 199]])

    env.run(menu, JobOptions(vendor="Tiny Place", city="pune"), upload=False)

    r = env.results()["Kaju Kolhapuri"]
    assert r["status"] == "ready" and r["source"] == "live"
    assert any(url == "https://www.zomato.com/pune/restaurants/maharashtrian"
               for url, _ in env.session.calls)


# ---- S2 TC2: the suggest query uses the spell-corrected dish name ----

def test_s2_tc2_suggest_query_uses_corrected_text(tmp_path):
    env = Env(tmp_path)
    env.session.add(SUGGEST, FakeResponse(200, suggest_json([])))
    menu = env.menu([["MASHRUM KOLHAPURI", "Mains", "", 199]])

    env.run(menu, JobOptions(vendor="Tiny Place", city="pune"), upload=False)

    suggest_calls = [params for url, params in env.session.calls if url == SUGGEST]
    assert any(p["q"] == "mushroom kolhapuri" for p in suggest_calls)


# ---- TC5: Zomato blocks every request ----

def test_blocked_continues(tmp_path):
    env = Env(tmp_path)
    env.session.add("https://www.zomato.com", FakeResponse(403, "forbidden"))
    menu = env.menu([["Farmhouse Pizza", "Pizza", "", 299], ["Cold Coffee", "Beverages", "", 120]])

    summary = env.run(menu, JobOptions(vendor="Some Cafe", city="pune"), upload=False)

    assert load_job(tmp_path / "job1")["status"] == "done"
    assert summary["needs_photo"] == 2
    assert all(r["status"] == "needs_photo" for r in env.results().values())
    assert any(e["type"] == "log" and "limiting requests" in e["message"] for e in env.events)


# ---- FIX2 TC2: an unexpected discovery exception (e.g. suggest raising ValueError) never aborts the job ----

def test_fix2_tc2_suggest_value_error_still_finishes_done(tmp_path, monkeypatch):
    env = Env(tmp_path)
    engine = env.engine()

    def boom(*args, **kwargs):
        raise ValueError("bad json shape")

    monkeypatch.setattr(engine.client, "suggest", boom)
    menu = env.menu([["Farmhouse Pizza", "Pizza", "", 299]])

    env.run(menu, JobOptions(vendor="Some Cafe", city="pune", live_search=False), engine=engine, upload=False)

    assert load_job(tmp_path / "job1")["status"] == "done"


# ---- TC5b: a block from one job does not disable live search forever ----

def test_blocked_expires_after_window(tmp_path, monkeypatch):
    env = Env(tmp_path)
    env.session.add("https://www.zomato.com", *([FakeResponse(403, "forbidden")] * 5), FakeResponse(404, "not found"))
    engine = env.engine()
    menu = env.menu([["Farmhouse Pizza", "Pizza", "", 299]])
    opts = JobOptions(vendor="Some Cafe", city="pune")

    env.run(menu, opts, engine=engine, upload=False)
    assert engine.blocked
    blocked_until = engine.blocked_until

    before = len(env.session.calls)
    env.run(menu, opts, job="job2", engine=engine, upload=False)
    assert not any(url.endswith("/order") for url, _ in env.session.calls[before:])
    assert any(e["type"] == "log" and "Live search paused" in e["message"] for e in env.events)
    assert engine.blocked_until == blocked_until  # not re-armed just by staying blocked

    monkeypatch.setattr(pipeline.time, "monotonic", lambda: blocked_until + 1)
    assert not engine.blocked
    job3_events: list[dict] = []
    env.run(menu, opts, job="job3", engine=engine, on_event=job3_events.append, upload=False)
    assert any(e["type"] == "stage" and e["stage"] == "search" for e in job3_events)


# ---- TC6: text-only mode when the vision model is unavailable ----

def test_text_only_mode(tmp_path):
    env = Env(tmp_path, vision=NullVision("offline"))
    env.pool("north-place", [("Paneer Butter", "Main Course", "veg", "pb1"),
                             ("Cheese Pizza", "Pizza", "veg", "cheese1")])
    env.pool("filler-place", [(f"{w} Sandwich", "Sandwiches", "veg", f"sw-{w}")
                              for w in ("Club", "Grilled", "Bombay", "Corn", "Aloo", "Mayo")])
    menu = env.menu([["Paneer Butter Masala", "Main Course", "", 250], ["Cheese Pizza", "Pizza", "", 199]])

    summary = env.run(menu, JobOptions(vendor="Some Cafe", live_search=False), upload=False)

    res = env.results()
    assert res["Paneer Butter Masala"]["status"] == "needs_photo"
    assert res["Paneer Butter Masala"]["reason"].startswith("Only uncertain matches")
    assert res["Cheese Pizza"]["status"] == "ready"
    assert summary["vision"] is False


# ---- TC7: resume does not re-download or re-upload finished dishes ----

def test_resume_skips_uploaded(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1"), ("Cold Coffee", "Drinks", "veg", "cc1")])
    menu = env.menu([["Cheese Pizza", "Pizza", "", 199], ["Cold Coffee", "Beverages", "", 120]])
    opts = JobOptions(vendor="Test Cafe", live_search=False)
    env.run(menu, opts)
    first = env.results()
    assert len(env.posts) == 2

    before = len(env.session.calls)
    summary = env.run(menu, opts)

    assert len(env.posts) == 2
    assert not any(url.startswith(IMG) and not url.endswith(THUMB_SUFFIX) for url, _ in env.session.calls[before:])
    assert summary["ready"] == 2
    assert {k: v["hosted_url"] for k, v in env.results().items()} == {k: v["hosted_url"] for k, v in first.items()}


# ---- TC8: two dishes never share one photo when an alternative exists ----

def test_distinct_images(tmp_path):
    env = Env(tmp_path)
    env.pool("place-a", [("Margherita Pizza", "Pizza", "veg", "shared"), ("Farmhouse Pizza", "Pizza", "veg", "shared")])
    env.pool("place-b", [("Spicy Margherita Pizza", "Pizza", "veg", "marg-b")])
    env.pool("place-c", [("Spicy Farmhouse Pizza", "Pizza", "veg", "farm-c")])
    menu = env.menu([["Margherita Pizza", "Pizza", "", 199], ["Farmhouse Pizza", "Pizza", "", 299]])

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    res = env.results()
    assert res["Margherita Pizza"]["zomato_img"] == IMG + "shared.jpg"
    assert res["Farmhouse Pizza"]["zomato_img"] == IMG + "farm-c.jpg"


# ---- TC9: unicode menu ----

def test_unicode_menu_end_to_end(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1")])
    menu = env.menu([["பனீர் டிக்கா", "Main Course", "घर का स्वाद", 220],
                     ["Cheese Pizza", "Pizza", "Loaded with cheese ₹ 199 😋", 199]])

    env.run(menu, JobOptions(vendor="Café Désí", live_search=False))

    assert load_job(tmp_path / "job1")["status"] == "done"
    res = env.results()
    assert res["பனீர் டிக்கா"]["status"] == "needs_photo"
    assert res["பனீர் டிக்கா"]["reason"] == "Name is not in English letters — add an English name"
    assert res["Cheese Pizza"]["status"] == "ready"
    for kind in ("UPLOAD_READY", "PENDING"):
        sheet(env, "Café Désí", kind)
    openpyxl.load_workbook(tmp_path / "job1" / f"{safe_vendor('Café Désí')}_Photo_Report.xlsx")
    assert sheet(env, "Café Désí", "UPLOAD_READY")["I2"].value == "Loaded with cheese ₹ 199 😋"


# ---- TC10: no ImgBB keys saved ----

def test_no_keys_marks_photo_rows_error(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1"), ("Cold Coffee", "Drinks", "veg", "cc1")])
    menu = env.menu([["Cheese Pizza", "Pizza", "", 199], ["Cold Coffee", "Beverages", "", 120],
                     ["Ceet-B Mastani", "Beverages", "", 90]])

    summary = env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), uploader=None)

    res = env.results()
    for name in ("Cheese Pizza", "Cold Coffee"):
        assert res[name]["status"] == "error" and res[name]["reason"] == NO_KEY_MESSAGE
    assert res["Ceet-B Mastani"]["status"] == "needs_photo"
    assert summary["errors"] == 2
    assert all((tmp_path / "job1" / f).is_file() for f in summary["files"].values())


# ---- TC11: cancel mid-run ----

def test_cancel_after_first_dish(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1"), ("Cold Coffee", "Drinks", "veg", "cc1"),
                             ("Farmhouse Pizza", "Pizza", "veg", "farm1")])
    menu = env.menu([["Cheese Pizza", "Pizza", "", 199], ["Cold Coffee", "Beverages", "", 120],
                     ["Farmhouse Pizza", "Pizza", "", 299]])
    cancel = threading.Event()

    def on_event(ev):
        env.events.append(ev)
        if ev["type"] == "dish":
            cancel.set()

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), on_event=on_event, cancel=cancel)

    assert load_job(tmp_path / "job1")["status"] == "cancelled"
    assert (tmp_path / "job1" / "results.json").is_file()
    assert len(env.posts) == 1
    assert not any(e["type"] == "done" for e in env.events)


# ---- TC12: menu without categories gets inferred ones ----

def test_blank_categories_inferred(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1"), ("Cold Coffee", "Drinks", "veg", "cc1")])
    menu = env.menu([["Cheese Pizza", 199], ["Cold Coffee", 120]], header=["Item Name", "Price"])

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    assert col(sheet(env, "Test Cafe", "UPLOAD_READY"), "H", 2) == ["Pizza", "Beverages"]


# ---- S1 TC7: a phonetic spelling is read before matching; the sheet keeps what the vendor wrote ----

def test_s1_tc7_misspelled_dish_is_read_and_kept(tmp_path):
    env = Env(tmp_path)
    env.pool("naan-house", [("Butter Naan", "Breads", "veg", "naan1")])
    menu = env.menu([["BATER NAN", "Roti", "", 40]])

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    assert env.results()["BATER NAN"]["status"] == "ready"
    assert col(sheet(env, "Test Cafe", "UPLOAD_READY"), "D", 1) == ["BATER NAN"]
    messages = [e["message"] for e in env.events if e["type"] == "log"]
    assert "Read 'BATER NAN' as 'butter naan'" in messages
    log_txt = (tmp_path / "job1" / "log.txt").read_text(encoding="utf-8")
    assert "Read 'BATER NAN' as 'butter naan'" in log_txt


# ---- TC13: a veg dish never takes a non-veg photo ----

def test_veg_dish_never_gets_nonveg_photo(tmp_path):
    env = Env(tmp_path)
    env.pool("momo-place", [("Momos", "Momos", "nonveg", "momo-nv"), ("Chicken Momos", "Momos", "nonveg", "cmomo")])
    menu = env.menu([["Veg Momos", "Momos", "", 120]])

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    r = env.results()["Veg Momos"]
    assert r["status"] == "needs_photo" and r["zomato_img"] == ""


# ---- TUNE1: calibrated vision veto (food_prob / head_margin) ----

def _candidate(res: str, url: str, tier: str) -> Candidate:
    item = PoolItem(0, "Dish", "Cat", "Delivery", "", "veg", url, res, res,
                     f"https://www.zomato.com/pune/{res}", "pune", res)
    return Candidate(item=item, score=MatchScore(tier=tier, text=0.9))


def _run_verify_group(env, cands: list[Candidate], head=None, monkeypatch=None,
                       food=None, margin=None, rows=None):
    engine = env.engine()
    for c in cands:
        vec = np.zeros(512, dtype=np.float32)
        vec[0] = 1.0
        engine.store.put_embedding(c.item.img_url, vec)
    if food is not None:
        monkeypatch.setattr(pipeline, "food_prob", lambda vision, mat: np.array(food))
    if margin is not None:
        monkeypatch.setattr(pipeline, "head_margin", lambda vision, mat, head: np.array(margin))
    q = DishQuery(raw="Dish", display="Dish", base_key="dish", variants=(frozenset({"DISH"}),),
                  head=head, diet="veg", proteins=frozenset(), is_combo=False)
    g = pipeline._Group("dish", q)
    g.cands = cands
    g.rows = rows or []
    run = pipeline._Run(engine, JobOptions(vendor="Test Cafe", live_search=False),
                         env.tmp / "job1", None, lambda e: None, lambda: None, None, {})
    with ThreadPoolExecutor(max_workers=2) as pool:
        run._verify_group(g, pool, Counter())
    return cands


def test_tune1_tc1_food_prob_040_kept(tmp_path, monkeypatch):
    env = Env(tmp_path)
    cands = [_candidate("r1", "https://img/1.jpg", "exact")]
    _run_verify_group(env, cands, monkeypatch=monkeypatch, food=[0.40], margin=[np.inf])

    assert cands[0].food_ok is True and cands[0].type_ok is True


def test_tune1_tc2_head_margin_hard_veto(tmp_path, monkeypatch):
    env = Env(tmp_path)

    kept = _candidate("r1", "https://img/kept.jpg", "exact")
    _run_verify_group(env, [kept], monkeypatch=monkeypatch, food=[0.9], margin=[-0.05])
    assert kept.type_ok is True

    rejected = _candidate("r2", "https://img/rejected.jpg", "exact")
    _run_verify_group(env, [rejected], monkeypatch=monkeypatch, food=[0.9], margin=[-0.12])
    assert rejected.type_ok is False


def test_tune1_tc3_soft_veto_removes_outlier_only(tmp_path, monkeypatch):
    env = Env(tmp_path)
    good = _candidate("r1", "https://img/good.jpg", "exact")
    outlier = _candidate("r2", "https://img/outlier.jpg", "exact")

    _run_verify_group(env, [good, outlier], monkeypatch=monkeypatch,
                       food=[0.9, 0.9], margin=[0.03, -0.05])

    assert good.type_ok is True
    assert outlier.type_ok is False


def test_tune1_tc4_close_tier_soft_veto_applies(tmp_path, monkeypatch):
    env = Env(tmp_path)
    close = _candidate("r1", "https://img/close.jpg", "close")

    _run_verify_group(env, [close], monkeypatch=monkeypatch, food=[0.9], margin=[-0.03])

    assert close.type_ok is False


def _verify_row() -> pipeline.DishResult:
    return pipeline.DishResult(row=1, sku="TC-001", name="Dish", category="Cat", description="", price=100)


def test_fix3_tc1_lone_close_candidate_needs_photo(tmp_path, monkeypatch):
    env = Env(tmp_path)
    row = _verify_row()
    _run_verify_group(env, [_candidate("r1", "https://img/close.jpg", "close")], monkeypatch=monkeypatch,
                       food=[0.9], margin=[0.05], rows=[row])

    assert row.zomato_img == ""
    assert row.reason == "Only uncertain matches found (best: Dish @ r1)"


def test_fix3_tc2_close_candidate_with_agreeing_photo_from_other_restaurant_kept(tmp_path, monkeypatch):
    env = Env(tmp_path)
    row = _verify_row()
    _run_verify_group(env, [_candidate("r1", "https://img/close1.jpg", "close"),
                            _candidate("r2", "https://img/close2.jpg", "close")],
                       monkeypatch=monkeypatch, food=[0.9, 0.9], margin=[0.05, 0.05], rows=[row])

    assert row.zomato_img in ("https://img/close1.jpg", "https://img/close2.jpg")
    assert row.tier == "close"


def test_fix3_close_candidates_from_same_restaurant_do_not_agree(tmp_path, monkeypatch):
    env = Env(tmp_path)
    row = _verify_row()
    first, second = _candidate("r1", "https://img/a.jpg", "close"), _candidate("r1", "https://img/b.jpg", "close")
    _run_verify_group(env, [first, second], monkeypatch=monkeypatch,
                       food=[0.9, 0.9], margin=[0.05, 0.05], rows=[row])

    assert row.zomato_img == ""


def test_fix3_lpp008_chain_menu_match_bypasses_close_gate_and_soft_veto(tmp_path, monkeypatch):
    env = Env(tmp_path)
    row = _verify_row()
    chain = _candidate("la-pinoz", "https://img/pack.jpg", "close")
    chain.source = "chain"
    _run_verify_group(env, [chain], monkeypatch=monkeypatch, food=[0.9], margin=[-0.05], rows=[row])

    assert row.zomato_img == "https://img/pack.jpg" and row.source == "chain"

    vetoed_row = _verify_row()
    vetoed = _candidate("la-pinoz", "https://img/wrong.jpg", "close")
    vetoed.source = "chain"
    _run_verify_group(env, [vetoed], monkeypatch=monkeypatch, food=[0.9], margin=[-0.12], rows=[vetoed_row])

    assert vetoed_row.zomato_img == ""


def test_tune1_tc5_food_prob_010_rejected_every_tier(tmp_path, monkeypatch):
    env = Env(tmp_path)
    for tier in ("exact", "strong", "close"):
        cand = _candidate("r1", f"https://img/{tier}.jpg", tier)
        _run_verify_group(env, [cand], monkeypatch=monkeypatch, food=[0.10], margin=[np.inf])
        assert cand.food_ok is False


# ---- TUNE3: brand guard (real RapidOCR reads visible text in photos) ----

def branded_jpeg(word: str, seed: int = 0) -> bytes:
    """A noisy but valid photo with `word` printed in a clear block of text."""
    blocks = np.random.default_rng(seed).integers(0, 256, (6, 6, 3), dtype=np.uint8)
    img = Image.fromarray(blocks).resize((420, 420), Image.NEAREST).convert("RGB")
    draw = ImageDraw.Draw(img)
    draw.rectangle([10, 160, 410, 260], fill=(255, 255, 255))
    draw.text((20, 175), word, fill=(0, 0, 0), font=ImageFont.truetype("arial.ttf", 48))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=95)
    return buf.getvalue()


def preset_image(env, name: str, content: bytes) -> str:
    """Register `content` for an image name before env.pool()/env.page() claim it."""
    url = IMG + name + ".jpg"
    env.session.add(url, FakeResponse(200, content=content))
    return url


def _photo_option(url: str, res: str) -> dict:
    return {"item_name": "Cheese Pizza", "restaurant": res, "restaurant_url": f"https://www.zomato.com/pune/{res}",
            "zomato_img": url, "tier": "exact", "confidence": 0.9, "source": "pool"}


def _photo_group_fixture(options: list[dict]) -> pipeline._Group:
    q = DishQuery(raw="Cheese Pizza", display="Cheese Pizza", base_key="cheese pizza",
                  variants=(frozenset({"CHEESE", "PIZZA"}),), head=None, diet="veg",
                  proteins=frozenset(), is_combo=False)
    g = pipeline._Group("cheese-pizza", q)
    r = pipeline.DishResult(row=1, sku="TC-001", name="Cheese Pizza", category="Pizza", description="", price=299)
    pipeline._apply_pick(r, options[0])
    r.alternates = options[1:]
    g.rows = [r]
    return g


def _run_photo_group(env, g, brand_words=frozenset()) -> pipeline._Run:
    run = pipeline._Run(env.engine(), JobOptions(vendor="Test Cafe", live_search=False),
                         env.tmp / "job1", None, env.events.append, lambda: None, None, {})
    run.brand_guard = BrandGuard(brand_words)
    (env.tmp / "job1" / "images").mkdir(parents=True, exist_ok=True)
    run._photo_group(g)
    return run


def test_tune3_tc3_branded_best_falls_back_to_clean_alternate(tmp_path):
    env = Env(tmp_path)
    branded_url = preset_image(env, "cheese-branded", branded_jpeg("ZOMATO"))
    clean_url = preset_image(env, "cheese-clean", noise_jpeg(9))
    g = _photo_group_fixture([_photo_option(branded_url, "res-a"), _photo_option(clean_url, "res-b")])

    _run_photo_group(env, g)

    r = g.rows[0]
    assert r.zomato_img == clean_url
    assert r.local_image == "images/TC-001.jpg"
    assert (env.tmp / "job1" / r.local_image).is_file()
    assert any(e.get("message") == "Skipped a photo showing zomato for Cheese Pizza" for e in env.events)


def test_tune3_tc4_all_branded_needs_photo(tmp_path):
    env = Env(tmp_path)
    b1 = preset_image(env, "branded-1", branded_jpeg("ZOMATO", seed=1))
    b2 = preset_image(env, "branded-2", branded_jpeg("ZOMATO", seed=2))
    g = _photo_group_fixture([_photo_option(b1, "res-a"), _photo_option(b2, "res-b")])

    _run_photo_group(env, g)

    r = g.rows[0]
    assert r.local_image == ""
    assert r.reason == "Only photos with zomato branding found — ask the owner for a photo"
    pipeline._settle(r, True)
    assert r.status == "needs_photo"


def test_tune3_tc5_vendor_own_brand_kept(tmp_path):
    env = Env(tmp_path)
    branded_url = preset_image(env, "special-pinoz", branded_jpeg("Pino'z"))
    env.pool("la-pinoz-0", [("Special Pizza", "Pizza", "veg", "special-pinoz")], res_name="La Pino'z Pizza")
    for i in range(1, 5):
        env.pool(f"la-pinoz-{i}", [("Other Dish", "Pizza", "veg", f"filler-{i}")], res_name="La Pino'z Pizza")
    menu = env.menu([["Special Pizza", "Pizza", "", 299]])

    env.run(menu, JobOptions(vendor="La Pino'z Pizza", live_search=False), upload=False)

    r = env.results()["Special Pizza"]
    assert r["source"] == "chain" and r["status"] == "ready" and r["zomato_img"] == branded_url


def test_fix2_tc4_fallback_never_reuses_another_dishs_photo(tmp_path):
    env = Env(tmp_path)
    used_url = preset_image(env, "used-elsewhere", noise_jpeg(11))
    alt_url = preset_image(env, "clean-alt", noise_jpeg(12))
    # best_url is never registered with the fake session, so its download fails (404 -> "http").
    best_url = IMG + "missing-best.jpg"

    q = DishQuery(raw="Cheese Pizza", display="Cheese Pizza", base_key="cheese pizza",
                  variants=(frozenset({"CHEESE", "PIZZA"}),), head=None, diet="veg",
                  proteins=frozenset(), is_combo=False)

    g_used = pipeline._Group("other-dish", q)
    r_used = pipeline.DishResult(row=1, sku="TC-001", name="Other Dish", category="Pizza", description="", price=299)
    pipeline._apply_pick(r_used, _photo_option(used_url, "res-b"))
    r_used.alternates = []
    g_used.rows = [r_used]

    g_fallback = pipeline._Group("cheese-pizza", q)
    r_fb = pipeline.DishResult(row=2, sku="TC-002", name="Cheese Pizza", category="Pizza", description="", price=299)
    pipeline._apply_pick(r_fb, _photo_option(best_url, "res-a"))
    r_fb.alternates = [_photo_option(used_url, "res-b"), _photo_option(alt_url, "res-c")]
    g_fallback.rows = [r_fb]

    run = pipeline._Run(env.engine(), JobOptions(vendor="Test Cafe", live_search=False),
                         env.tmp / "job1", None, env.events.append, lambda: None, None, {})
    run.brand_guard = BrandGuard(frozenset())
    (env.tmp / "job1" / "images").mkdir(parents=True, exist_ok=True)
    # Mimics what the verify stage records for its best pick before photos() ever runs.
    run.used_images[used_url] = g_used.key

    run._photo_group(g_used)
    run._photo_group(g_fallback)

    assert g_used.rows[0].zomato_img == used_url
    assert g_fallback.rows[0].zomato_img == alt_url
    assert g_fallback.rows[0].zomato_img != g_used.rows[0].zomato_img


def test_tune3_tc6_ocr_import_failure_disables_guard(tmp_path, monkeypatch):
    class BrokenRapidOCR:
        def __init__(self, *a, **kw):
            raise RuntimeError("models unavailable")

    broken = type(sys)("rapidocr")
    broken.RapidOCR = BrokenRapidOCR
    monkeypatch.setitem(sys.modules, "rapidocr", broken)
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1")])
    menu = env.menu([["Cheese Pizza", "Pizza", "", 199]])

    summary = env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    assert summary["ready"] == 1
    res = env.results()["Cheese Pizza"]
    assert res["status"] == "ready" and res["zomato_img"]


def test_fix3_tc3_logo_check_unavailable_drops_big_chain_photos(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "rapidocr", None)
    env = Env(tmp_path)
    env.pool("wowza-0", [("Cheese Pizza", "Pizza", "veg", "cheese-wowza")], res_name="Wowza Pizza")
    for i in range(1, 5):
        env.pool(f"wowza-{i}", [("Other Dish", "Pizza", "veg", f"wowza-filler-{i}")], res_name="Wowza Pizza")
    env.pool("corner-cafe", [("Cold Coffee", "Drinks", "veg", "cc-corner")])
    menu = env.menu([["Cheese Pizza", "Pizza", "", 199], ["Cold Coffee", "Beverages", "", 99]])

    summary = env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    assert any(w.startswith("Logo check unavailable") for w in summary["warnings"])
    res = env.results()
    assert res["Cheese Pizza"]["status"] == "needs_photo" and res["Cheese Pizza"]["local_image"] == ""
    assert res["Cheese Pizza"]["reason"] == ("Only photos from Wowza Pizza found and the logo check is unavailable"
                                             " — ask the owner for a photo")
    assert res["Cold Coffee"]["status"] == "ready"
    report = openpyxl.load_workbook(env.tmp / "job1" / f"{safe_vendor('Test Cafe')}_Photo_Report.xlsx")["Summary"]
    assert {row[0].value: row[1].value for row in report.iter_rows()}["Logo check"] == "Unavailable"


def test_fix3_tc5_generic_same_name_pool_restaurants_are_not_a_chain(tmp_path):
    env = Env(tmp_path)
    for i in range(3):
        env.pool(f"spice-kitchen-{i}", [("Paneer Tikka", "Starters", "veg", f"pt-{i}")], res_name="Spice Kitchen")
    env.session.add(SUGGEST, FakeResponse(200, suggest_json([])))
    menu = env.menu([["Paneer Tikka", "Starters", "", 199]])

    env.run(menu, JobOptions(vendor="Spice Kitchen", live_search=False), upload=False)

    assert env.results()["Paneer Tikka"]["source"] != "chain"
    assert any(e.get("message") == "Vendor not found on Zomato; using photos from other restaurants"
               for e in env.events)


def test_fix3_own_url_menu_always_in_source_set(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "MAX_SOURCE_OUTLETS", 1)
    env = Env(tmp_path)
    env.pool("my-cafe-kothrud", [("Garden Special", "Pizza", "veg", "garden-other"),
                                 ("Cold Coffee", "Drinks", "veg", "cc-other")], res_name="My Cafe")
    env.page("my-cafe", [("Garden Special", "Pizza", "veg", "garden-own")], res_name="My Cafe")
    menu = env.menu([["Garden Special", "Pizza", "", 299]])

    env.run(menu, JobOptions(vendor="My Cafe", zomato_url="https://www.zomato.com/pune/my-cafe/order",
                             live_search=False), upload=False)

    r = env.results()["Garden Special"]
    assert r["source"] == "own" and r["zomato_img"] == IMG + "garden-own.jpg"


def test_uifix_tc2_low_disk_fails_before_any_work(tmp_path, monkeypatch):
    class FakeUsage:
        free = 0.2 * 1_000_000_000

    monkeypatch.setattr(pipeline.shutil, "disk_usage", lambda path: FakeUsage())
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1")])
    menu = env.menu([["Cheese Pizza", "Pizza", "", 199]])

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    job = load_job(tmp_path / "job1")
    assert job["status"] == "failed"
    assert "disk is almost full" in job["error"]
    assert not (tmp_path / "job1" / "results.json").exists()
    assert not list((tmp_path / "job1").glob("*_SmartBiz_*.xlsx"))


def test_fix1_tc5_low_disk_keeps_menu_and_resumes(tmp_path, monkeypatch):
    class FakeUsage:
        free = 0.2 * 1_000_000_000

    real_disk_usage = pipeline.shutil.disk_usage
    monkeypatch.setattr(pipeline.shutil, "disk_usage", lambda path: FakeUsage())
    env = Env(tmp_path)
    env.pool("pizza-place", [("Cheese Pizza", "Pizza", "veg", "cheese1")])
    menu = env.menu([["Cheese Pizza", "Pizza", "", 199]])
    opts = JobOptions(vendor="Test Cafe", live_search=False)

    env.run(menu, opts, upload=False)

    job = load_job(tmp_path / "job1")
    assert job["status"] == "failed"
    assert job["menu_file"] == "menu.xlsx"
    assert (tmp_path / "job1" / "menu.xlsx").is_file()
    errors = [e for e in env.events if e["type"] == "error"]
    assert errors and "disk is almost full" in errors[0]["message"]
    created = job["created"]

    monkeypatch.setattr(pipeline.shutil, "disk_usage", real_disk_usage)
    summary = env.run(tmp_path / "job1" / job["menu_file"], opts, upload=False)

    job = load_job(tmp_path / "job1")
    assert job["status"] == "done"
    assert job["created"] == created
    assert summary["total"] == 1


# ---- FIX2 TC7: atomic_write_text survives transient PermissionError (OneDrive/antivirus lock) ----

def test_fix2_tc7_atomic_write_text_retries_permission_error(tmp_path, monkeypatch):
    path = tmp_path / "job.json"
    import os as os_mod

    attempts = {"n": 0}
    real_replace = os_mod.replace

    def flaky_replace(src, dst):
        attempts["n"] += 1
        if attempts["n"] <= 2:
            raise PermissionError(13, "Permission denied", str(dst))
        return real_replace(src, dst)

    monkeypatch.setattr(os_mod, "replace", flaky_replace)
    atomic_write_text(path, "hello")

    assert attempts["n"] == 3
    assert path.read_text(encoding="utf-8") == "hello"


# ---- A5: look-alike rungs, Swiggy search, review status, search/approve helpers ----

LOOKALIKE_DISH = "Green Chilli Paneer Pizza (Single Cheese)"
SWIGGY_FIXTURE = Path(__file__).parent / "fixtures" / "swiggy" / "dish_search_mushroom_uttapam.json"


def _lookalike_job(env) -> dict:
    env.pool("pizza-place", [("Paneer Pizza", "Pizza", "veg", "paneer-pizza")])
    menu = env.menu([[LOOKALIKE_DISH, "Pizza", "", 120]])
    return env.run(menu, JobOptions(vendor="Test Cafe", live_search=False))


def test_a5_tc1_lookalike_photo_goes_to_review(tmp_path):
    env = Env(tmp_path)

    summary = _lookalike_job(env)

    r = env.results()[LOOKALIKE_DISH]
    assert r["status"] == "review" and r["lookalike"] is True and r["search_name"] == "Paneer Pizza"
    assert r["photo_names"] == ["Green Chilli Paneer Pizza", "Paneer Pizza"]
    assert r["reason"] == "Look-alike photo (Paneer Pizza) — approve it or search again"
    assert r["local_image"] and (tmp_path / "job1" / r["local_image"]).is_file()
    assert r["hosted_url"] == "" and env.posts == []
    assert col(sheet(env, "Test Cafe", "UPLOAD_READY"), "D", 1) == [None]
    assert col(sheet(env, "Test Cafe", "PENDING"), "D", 1) == [LOOKALIKE_DISH]
    assert summary["review"] == 1
    assert [e["result"]["status"] for e in env.events if e["type"] == "dish"] == ["review"]


def test_a5_tc2_photo_names_column_is_searched(tmp_path):
    env = Env(tmp_path)
    env.pool("grill-place", [("Chicken Lollipop", "Starters", "nonveg", "lollipop")])
    # retrieval skips tokens found in over a quarter of a tiny pool, so pad it
    env.pool("filler-place", [("Cheese Pizza", "Pizza", "veg", "cheese1"), ("Cold Coffee", "Drinks", "veg", "cc1"),
                              ("Mango Shake", "Shakes", "veg", "mango1")])
    menu = env.menu([["Tandoori Lollipop", "Starters", "", 220, "Non-Veg", "Chicken Lollipop"]],
                    header=[*HEADER, "Diet", "Photo Names"])

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    r = env.results()["Tandoori Lollipop"]
    assert r["status"] == "review" and r["search_name"] == "Chicken Lollipop" and r["diet"] == "nonveg"
    assert r["photo_names"] == ["Chicken Lollipop"]


def test_a5_tc3_own_name_match_wins(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Paneer Pizza", "Pizza", "veg", "paneer-pizza")])
    menu = env.menu([["Paneer Pizza", "Pizza", "", 199]])

    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False), upload=False)

    r = env.results()["Paneer Pizza"]
    assert r["status"] == "ready" and r["lookalike"] is False and r["search_name"] == ""


def _route_swiggy(env, status: int = 200) -> None:
    text = SWIGGY_FIXTURE.read_text(encoding="utf-8")
    env.session.add(SEARCH_URL, FakeResponse(status, text if status == 200 else "denied"))
    for menu in parse_dish_search(json.loads(text), "pune"):
        for it in menu.items:
            jpeg = noise_jpeg(zlib.crc32(it.img_url.encode()))
            env.session.add(thumb_url(it.img_url), FakeResponse(200, content=jpeg))
            env.session.add(full_url(it.img_url), FakeResponse(200, content=jpeg))


def test_a5_tc4_swiggy_search_finds_photo(tmp_path):
    env = Env(tmp_path)
    _route_swiggy(env)
    env.session.add(SUGGEST, FakeResponse(200, suggest_json([])))
    menu = env.menu([["Masala Uttapam", "South Indian", "", 90]])
    opts = JobOptions(vendor="Test Cafe", city="pune")

    env.run(menu, opts, upload=False)

    assert env.results()["Masala Uttapam"]["status"] == "needs_photo"
    assert not any(url == SEARCH_URL for url, _ in env.session.calls)

    engine = Engine(env.store, env.vision, env.http, swiggy=SwiggyClient(env.http))
    env.run(menu, opts, job="job2", engine=engine, upload=False)

    r = env.results("job2")["Masala Uttapam"]
    assert r["status"] == "ready" and r["source"] == "swiggy"


def test_a5_tc5_swiggy_blocked_continues(tmp_path):
    env = Env(tmp_path)
    http = Http(session=env.session, min_interval=0.0, max_retries=0, sleep=lambda s: None)
    _route_swiggy(env, status=403)
    env.session.add(SUGGEST, FakeResponse(200, suggest_json([])))
    menu = env.menu([["Masala Uttapam", "South Indian", "", 90], ["Onion Dosa", "South Indian", "", 80]])
    engine = Engine(env.store, env.vision, http, swiggy=SwiggyClient(http))

    summary = env.run(menu, JobOptions(vendor="Test Cafe", city="pune"), engine=engine, upload=False)

    assert load_job(tmp_path / "job1")["status"] == "done"
    assert summary["warnings"].count(SWIGGY_BLOCKED_MESSAGE) == 1
    assert sum(url == SEARCH_URL for url, _ in env.session.calls) == 1
    assert env.results()["Masala Uttapam"]["status"] == "needs_photo"


def test_a5_tc6_approve_uploads_and_readies(tmp_path):
    env = Env(tmp_path)
    _lookalike_job(env)
    job_dir, engine = tmp_path / "job1", env.engine()

    [r] = approve_dishes(engine, job_dir, ["TC-001"], env.uploader)

    assert r.status == "ready" and r.approved is True and r.hosted_url.startswith("https://i.ibb.co/")
    assert env.results()[LOOKALIKE_DISH]["hosted_url"] == r.hosted_url
    assert load_job(job_dir)["sheets_stale"] is True
    rebuild_outputs(engine, job_dir)
    assert col(sheet(env, "Test Cafe", "UPLOAD_READY"), "D", 1) == [LOOKALIKE_DISH]

    before = env.results()[LOOKALIKE_DISH]
    [again] = approve_dishes(engine, job_dir, ["TC-001"], env.uploader)
    assert again.status == "ready" and len(env.posts) == 1
    assert env.results()[LOOKALIKE_DISH] == before

    with pytest.raises(ValueError, match="Unknown SKU TC-999"):
        approve_dishes(engine, job_dir, ["TC-999"], env.uploader)


def test_a5_tc7_search_photos_then_choose(tmp_path):
    env = Env(tmp_path)
    env.pool("thali-place", [("Veg Thali", "Thali", "veg", "veg-thali")])
    env.pool("pizza-place", [("Paneer Pizza", "Pizza", "veg", "paneer-pizza")])
    menu = env.menu([["Shubham All-in-One Thali", "Thali", "", 250]])
    env.run(menu, JobOptions(vendor="Test Cafe", live_search=False))
    job_dir, engine = tmp_path / "job1", env.engine()
    first = env.results()["Shubham All-in-One Thali"]
    assert first["status"] in ("review", "needs_photo")

    r = search_photos(engine, job_dir, "TC-001", "Veg Thali")

    assert r.search_name == "Veg Thali"
    assert any(a["item_name"] == "Veg Thali" and a["rung"] >= 1 for a in r.alternates)
    assert all(a["confidence"] > 0 for a in r.alternates)
    assert r.zomato_img == first["zomato_img"] and r.status == first["status"]
    assert env.results()["Shubham All-in-One Thali"]["alternates"] == r.alternates
    with pytest.raises(ValueError, match="Type a dish name to search for"):
        search_photos(engine, job_dir, "TC-001", "  ")

    chosen = choose_alternate(engine, job_dir, "TC-001", 0, env.uploader)

    assert chosen.status == "ready" and chosen.approved is True and chosen.lookalike is True


# ---- A5 round 1 ----

def test_a5_r1_rerun_keeps_approved_lookalike(tmp_path):
    env = Env(tmp_path)
    _lookalike_job(env)
    approve_dishes(env.engine(), tmp_path / "job1", ["TC-001"], env.uploader)
    hosted = env.results()[LOOKALIKE_DISH]["hosted_url"]

    summary = env.run(tmp_path / "menu.xlsx", JobOptions(vendor="Test Cafe", live_search=False))

    r = env.results()[LOOKALIKE_DISH]
    assert r["status"] == "ready" and r["approved"] is True and r["reason"] == ""
    assert r["lookalike"] is True and r["search_name"] == "Paneer Pizza"
    assert r["hosted_url"] == hosted and len(env.posts) == 1 and summary["ready"] == 1


def test_a5_r1_choose_back_keeps_lookalike_state(tmp_path):
    env = Env(tmp_path)
    env.pool("pizza-place", [("Paneer Pizza", "Pizza", "veg", "paneer-pizza")])
    env.run(env.menu([["Paneer Pizza", "Pizza", "", 199]]), JobOptions(vendor="Test Cafe", live_search=False),
            upload=False)
    job_dir, engine = tmp_path / "job1", env.engine()
    search_photos(engine, job_dir, "TC-001", "Paneer Pizza")

    picked = choose_alternate(engine, job_dir, "TC-001", 0, None)
    assert picked.lookalike is True and picked.search_name == "Paneer Pizza"

    back = choose_alternate(engine, job_dir, "TC-001", 0, None)
    assert back.lookalike is False and back.search_name == ""
    assert back.alternates[0]["rung"] == 1

    again = choose_alternate(engine, job_dir, "TC-001", 0, None)
    assert again.lookalike is True


def test_a5_r1_approve_missing_photo_file(tmp_path):
    env = Env(tmp_path)
    _lookalike_job(env)
    (tmp_path / "job1" / env.results()[LOOKALIKE_DISH]["local_image"]).unlink()

    with pytest.raises(ValueError, match="Photo file missing for TC-001; search again"):
        approve_dishes(env.engine(), tmp_path / "job1", ["TC-001"], env.uploader)
    assert env.posts == []
