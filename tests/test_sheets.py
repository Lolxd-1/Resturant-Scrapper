"""Tests for studio.sheets (SmartBiz output builder)."""
from __future__ import annotations

import time
from pathlib import Path

import openpyxl
import pytest

from studio.models import DishResult
from studio.sheets import STATUS_LABELS, build_outputs, category_for, safe_vendor

REPO_ROOT = Path(__file__).resolve().parents[1]
REAL_TEMPLATE = REPO_ROOT / "data" / "smartbiz_template.xlsx"

BULK_HEADER = [
    "SKU ID", "Variant ID", "Custom SKU", "Product Name", "MRP", "Selling Price",
    "Business Category", "Product Category", "Description", "Variant Relationship",
    "Size", "Colour Code", "Colour Name", "Best Seller", "HSN Code",
    "Product Image1", "Product Image2", "Product Image3", "Product Image4",
    "Product Image5", "Product Image6",
]


def make_template(path: Path, validation_rows: int = 20) -> Path:
    """A small stand-in for data/smartbiz_template.xlsx with the same sheet layout."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "important_instructions"
    ws1.append(["Read me first"])

    ws2 = wb.create_sheet("bulk_upload_template")
    ws2.append(BULK_HEADER)
    for _ in range(validation_rows):
        ws2.append([None] * len(BULK_HEADER))

    ws3 = wb.create_sheet("DataSheet")
    ws3.append(["lookup data"])
    ws3.sheet_state = "hidden"

    wb.save(path)
    return path


def mk(row, sku, name, category="Pizza", description="", price=100.0, status="ready", **kw) -> DishResult:
    return DishResult(
        row=row, sku=sku, name=name, category=category, description=description,
        price=price, status=status, **kw,
    )


def read_col(ws, col_letter, min_row, max_row):
    return [ws[f"{col_letter}{r}"].value for r in range(min_row, max_row + 1)]


# ---- TC1: categories written to column H, never the old hardcoded value ----

def test_tc1_categories_in_column_h(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-001", "Margherita", category="Pizza"),
        mk(2, "SKU-002", "Cola", category="Beverages"),
        mk(3, "SKU-003", "Thali", category="Main Course"),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    wb = openpyxl.load_workbook(paths["ready"])
    ws = wb["bulk_upload_template"]
    cats = read_col(ws, "H", 2, 4)
    assert cats == ["Pizza", "Beverages", "Main Course"]
    assert "Other Food and Grocery" not in cats


# ---- TC2: rows written contiguously; ready/pending split correctly ----

def test_tc2_contiguous_rows(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-001", "Dish1"),
        mk(2, "SKU-002", "Dish2", status="needs_photo"),
        mk(3, "SKU-003", "Dish3"),
        mk(4, "SKU-004", "Dish4", status="needs_photo"),
        mk(5, "SKU-005", "Dish5"),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    ready_ws = openpyxl.load_workbook(paths["ready"])["bulk_upload_template"]
    assert read_col(ready_ws, "C", 2, 4) == ["SKU-001", "SKU-003", "SKU-005"]
    assert ready_ws["C5"].value is None

    pending_ws = openpyxl.load_workbook(paths["pending"])["bulk_upload_template"]
    assert read_col(pending_ws, "C", 2, 3) == ["SKU-002", "SKU-004"]


# ---- TC3: needs_price with a photo already found goes to PENDING with P set, E/F blank ----

def test_tc3_needs_price_with_photo(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-001", "Dish1", price=None, status="needs_price", hosted_url="https://i.ibb.co/x/a.jpg"),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    pending_ws = openpyxl.load_workbook(paths["pending"])["bulk_upload_template"]
    assert pending_ws["P2"].value == "https://i.ibb.co/x/a.jpg"
    assert pending_ws["E2"].value is None
    assert pending_ws["F2"].value is None

    ready_ws = openpyxl.load_workbook(paths["ready"])["bulk_upload_template"]
    assert ready_ws["C2"].value is None


# ---- TC1b: PENDING never inherits READY's values (openpyxl.cell(value=None) does not clear) ----

def test_tc1b_pending_does_not_inherit_ready_values(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-A", "DishA", status="ready", hosted_url="https://i.ibb.co/a.jpg", price=100),
        mk(2, "SKU-B", "DishB", status="needs_photo", price=50, description=""),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    ready_ws = openpyxl.load_workbook(paths["ready"])["bulk_upload_template"]
    assert ready_ws["D2"].value == "DishA"
    assert ready_ws["C3"].value is None

    pending_ws = openpyxl.load_workbook(paths["pending"])["bulk_upload_template"]
    assert pending_ws["D2"].value == "DishB"
    assert pending_ws["P2"].value is None  # no hosted_url; must not show DishA's photo
    assert pending_ws["I2"].value is None  # no description; must not show DishA's row leftovers


# ---- TC4: field length limits ----

def test_tc4_limits(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-001", "N" * 250, category="C" * 150, description="D" * 3000),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    ws = openpyxl.load_workbook(paths["ready"])["bulk_upload_template"]
    assert len(ws["D2"].value) == 200
    assert len(ws["I2"].value) == 2000
    assert len(ws["H2"].value) == 100


# ---- TC5: template structure untouched ----

def test_tc5_template_intact(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [mk(1, "SKU-001", "Dish1")]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    src_wb = openpyxl.load_workbook(template)
    out_wb = openpyxl.load_workbook(paths["ready"])
    assert out_wb.sheetnames == ["important_instructions", "bulk_upload_template", "DataSheet"]

    src_header = [c.value for c in src_wb["bulk_upload_template"][1]]
    out_header = [c.value for c in out_wb["bulk_upload_template"][1]]
    assert out_header == src_header
    assert out_wb["DataSheet"].sheet_state == "hidden"


# ---- TC6: unicode round-trips; illegal XML control chars are stripped ----

def test_tc6_unicode_and_illegal_chars(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-001", "पнीर टिक्का"),
        mk(2, "SKU-002", "Crème ₹ \U0001f355"),
        mk(3, "SKU-003", "Bad\x07Name"),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    ws = openpyxl.load_workbook(paths["ready"])["bulk_upload_template"]
    assert ws["D2"].value == "पнीर टिक्का"
    assert ws["D3"].value == "Crème ₹ \U0001f355"
    assert ws["D4"].value == "BadName"
    assert "\x07" not in ws["D4"].value


# ---- TC7: report workbook ----

def test_tc7_report_contents(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-001", "Dish1", status="ready", hosted_url="https://i.ibb.co/1.jpg", confidence=0.91),
        mk(2, "SKU-002", "Dish2", status="needs_photo", reason="no match found"),
        mk(3, "SKU-003", "Dish3", status="needs_price", price=None, hosted_url="https://i.ibb.co/3.jpg"),
        mk(4, "SKU-004", "Dish4", status="error", reason="bad row"),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    wb = openpyxl.load_workbook(paths["report"])
    dishes = wb["Dishes"]
    assert [dishes.cell(row=r, column=1).value for r in range(2, 6)] == [
        "SKU-001", "SKU-002", "SKU-003", "SKU-004",
    ]
    assert [dishes.cell(row=r, column=5).value for r in range(2, 6)] == [
        "Ready", "Needs owner photo", "Needs price", "Error",
    ]

    checklist = wb["Owner photo checklist"]
    checklist_skus = [checklist.cell(row=r, column=1).value for r in range(2, checklist.max_row + 1)]
    assert checklist_skus == ["SKU-002"]

    summary = wb["Summary"]
    summary_vals = {row[0].value: row[1].value for row in summary.iter_rows(min_row=1, max_row=summary.max_row)}
    # ready(1) + needs_price-with-photo(1) = 2 of 4 total -> 50%
    assert summary_vals["Photo coverage"] == "50%"
    assert summary_vals["Dishes in menu"] == 4
    assert summary_vals["Ready to upload"] == 1
    assert summary_vals["Need owner photo"] == 1
    assert summary_vals["Need price"] == 1
    assert summary_vals["Errors"] == 1
    assert summary_vals["Visual check"] == "On"


# ---- TC8: performance against the real template ----

@pytest.mark.skipif(not REAL_TEMPLATE.exists(), reason="real SmartBiz template not present")
def test_tc8_speed_real_template(tmp_path):
    results = [
        mk(i, f"SKU-{i:03d}", f"Dish {i}", status="ready" if i % 3 else "needs_photo")
        for i in range(1, 151)
    ]
    t0 = time.time()
    build_outputs(results, REAL_TEMPLATE, tmp_path / "out", vendor="Acme")
    assert time.time() - t0 < 60


# ---- TC9: price formatting ----

def test_tc9_price_formatting(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [
        mk(1, "SKU-001", "Dish1", price=120.0),
        mk(2, "SKU-002", "Dish2", price=99.5),
    ]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    ws = openpyxl.load_workbook(paths["ready"])["bulk_upload_template"]
    assert ws["E2"].value == 120 and isinstance(ws["E2"].value, int)
    assert ws["F2"].value == 120
    assert ws["E3"].value == 99.5
    assert ws["F3"].value == 99.5


# ---- TC10: blank category falls back to "Menu" ----

def test_tc10_blank_category_falls_back(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [mk(1, "SKU-001", "Dish1", category="  ")]
    paths = build_outputs(results, template, tmp_path / "out", vendor="Acme")

    ws = openpyxl.load_workbook(paths["ready"])["bulk_upload_template"]
    assert ws["H2"].value == "Menu"
    assert category_for(results[0]) == "Menu"


# ---- safe_vendor ----

def test_safe_vendor():
    assert safe_vendor("Joe's Pizza & Grill!!") == "Joe_s_Pizza_Grill"
    assert safe_vendor("") == "Vendor"
    assert safe_vendor("a" * 60) == "a" * 40


def test_build_outputs_file_names(tmp_path):
    template = make_template(tmp_path / "template.xlsx")
    results = [mk(1, "SKU-001", "Dish1")]
    paths = build_outputs(results, template, tmp_path / "out", vendor="My Vendor!!")

    assert paths["ready"].name == "My_Vendor_SmartBiz_UPLOAD_READY.xlsx"
    assert paths["pending"].name == "My_Vendor_SmartBiz_PENDING.xlsx"
    assert paths["report"].name == "My_Vendor_Photo_Report.xlsx"
    assert set(STATUS_LABELS) == {"ready", "needs_photo", "needs_price", "error"}
